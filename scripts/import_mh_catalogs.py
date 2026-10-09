#!/usr/bin/env python3
"""Import the user-supplied MH catalog workbook; no network or credentials.

Maintenance only: requires openpyxl. Run with the official XLSX as a positional
argument. Leading zeroes and string codes are preserved. Workbook modification
metadata describes the file and must not be interpreted as legal validity.
"""

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import unicodedata


ROOT = Path(__file__).resolve().parent.parent
INTEGER_CATALOGS = {
    "CAT-003", "CAT-004", "CAT-005", "CAT-007", "CAT-010", "CAT-011",
    "CAT-014", "CAT-016", "CAT-021", "CAT-024", "CAT-026", "CAT-029",
    "CAT-030", "CAT-032",
}
EXPECTED_DISTRICTS = [12, 13, 16, 33, 22, 19, 16, 22, 9, 13, 23, 20, 26, 18]
EXPECTED_MUNICIPALITIES = [3, 4, 4, 3, 6, 5, 2, 3, 2, 2, 3, 3, 2, 2]


def folded(value):
    return "".join(char for char in unicodedata.normalize("NFD", value.upper())
                   if not unicodedata.combining(char))


def import_workbook(source):
    from openpyxl import load_workbook

    source = Path(source)
    book = load_workbook(source, data_only=True)
    if book.sheetnames != ["Hoja1"]:
        raise ValueError("Revisa el formato del libro: se esperaba una sola hoja Hoja1.")
    sheet = book.active
    version_title = str(sheet["A5"].value or "").strip()
    version = re.search(r"\bV\s+(\d+\.\d+)\b", version_title)
    if not version:
        raise ValueError("No se encontró la versión del catálogo en A5.")
    catalogs = {}
    active = None
    group = None
    for row in sheet.iter_rows():
        cell, label_cell = row[:2]
        value, label = cell.value, label_cell.value
        heading = re.fullmatch(r"(CAT-\d{3})\s+(.+)", str(value).strip()) if value is not None else None
        if heading:
            catalog_id, name = heading.groups()
            if catalog_id in catalogs:
                raise ValueError(f"Encabezado duplicado: {catalog_id}.")
            active = catalogs[catalog_id] = {
                "name": name.strip(),
                "code_type": "integer" if catalog_id in INTEGER_CATALOGS else "string",
                "entries": [],
                "notes": [],
            }
            group = None
            continue
        if active is None or value is None and label is None:
            continue
        if str(value).strip().upper() == "CÓDIGO":
            continue
        if value is None or label is None:
            text = str(label if value is None else value).strip()
            active["notes"].append({"source_row": cell.row, "text": text})
            if catalog_id in ("CAT-015", "CAT-019"):
                group = text
            continue
        if isinstance(value, bool) or isinstance(value, float) and not value.is_integer():
            raise ValueError(f"Código numérico inesperado en A{cell.row}.")
        code = str(int(value)) if isinstance(value, float) else str(value).strip()
        # The supplied workbook keeps fixed-width codes as strings when needed;
        # numeric cells use only codes whose decimal text already has that width.
        if isinstance(value, int) and catalog_id not in INTEGER_CATALOGS:
            width = 5 if catalog_id == "CAT-019" else 2
            if len(code) != width:
                raise ValueError(f"Código ambiguo sin ceros iniciales en A{cell.row}: {code}.")
        entry = {"code": code, "label": str(label).strip(), "source_row": cell.row}
        if group:
            entry["group"] = group
        active["entries"].append(entry)
    if set(catalogs) != {f"CAT-{number:03}" for number in range(1, 34)}:
        raise ValueError("Se esperaban los 33 catálogos CAT-001 a CAT-033.")

    # CAT-008 contains each department's legacy districts in consecutive blocks.
    # Codes restart at 01; this avoids losing repeated district codes in a dict.
    department = 0
    for entry in catalogs["CAT-008"]["entries"]:
        if entry["code"] == "00":
            entry["department"] = "00"
            continue
        if entry["code"] == "01":
            department += 1
        entry["department"] = f"{department:02}"
    counts = Counter(entry["department"] for entry in catalogs["CAT-008"]["entries"])
    if [counts[f"{number:02}"] for number in range(1, 15)] != EXPECTED_DISTRICTS:
        raise ValueError("La agrupación departamental de CAT-008 cambió; revisa la fuente.")

    # CAT-013 labels contain their department name, unlike the repeated numeric
    # codes. Derive parents from exact names, without guessing code ranges.
    departments = {folded(entry["label"]): entry["code"]
                   for entry in catalogs["CAT-012"]["entries"] if entry["code"] != "00"}
    for entry in catalogs["CAT-013"]["entries"]:
        if entry["code"] == "00":
            entry["department"] = "00"
            continue
        matches = [code for name, code in departments.items()
                   if folded(entry["label"]).startswith(name + " ")]
        if len(matches) != 1:
            raise ValueError(f"No se pudo resolver departamento para {entry['label']}.")
        entry["department"] = matches[0]
    counts = Counter(entry["department"] for entry in catalogs["CAT-013"]["entries"])
    if [counts[f"{number:02}"] for number in range(1, 15)] != EXPECTED_MUNICIPALITIES:
        raise ValueError("La agrupación departamental de CAT-013 cambió; revisa la fuente.")

    for catalog_id, catalog in catalogs.items():
        if not catalog["entries"]:
            raise ValueError(f"Catálogo vacío: {catalog_id}.")
        seen = set()
        for entry in catalog["entries"]:
            identity = (entry["code"], entry.get("department"), entry.get("group"))
            if identity in seen:
                raise ValueError(f"Código duplicado sin contexto: {catalog_id} {entry['code']}.")
            seen.add(identity)
    metadata = {
        "publisher": "Ministerio de Hacienda de El Salvador",
        "source_origin": "Documento compartido por el usuario desde factura.gob.sv",
        "source_filename": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_title": version_title,
        "version": version.group(1),
        "workbook_created_at": book.properties.created.isoformat() if book.properties.created else None,
        "workbook_modified_at": book.properties.modified.isoformat() if book.properties.modified else None,
        "effective_date": None,
        "date_note": "Las fechas del libro son metadatos del archivo, no fechas de vigencia normativa.",
        "territorial_note": "CAT-013 contiene 44 municipios; CAT-008 contiene 262 distritos. Ambos incluyen 00 para extranjeros. El libro no declara la relación municipio-distrito.",
    }
    return {"metadata": metadata, "catalogs": catalogs}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "app/schemas/mh-catalogs.json")
    args = parser.parse_args()
    dataset = import_workbook(args.workbook)
    encoded = (json.dumps(dataset, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(encoded)
    provenance = {"dataset_filename": args.output.name,
                  "dataset_sha256": hashlib.sha256(encoded).hexdigest(),
                  **dataset["metadata"]}
    provenance_path = args.output.with_name(args.output.stem + "-provenance.json")
    provenance_path.write_text(json.dumps(provenance, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Importados {len(dataset['catalogs'])} catálogos y "
          f"{sum(len(c['entries']) for c in dataset['catalogs'].values())} registros.")
    print(f"SHA-256 de la fuente: {dataset['metadata']['source_sha256']}")
    print(f"SHA-256 del dataset: {provenance['dataset_sha256']}")


if __name__ == "__main__":
    main()
